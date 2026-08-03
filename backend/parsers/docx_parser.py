"""DOCX text extraction using python-docx.

Extracts text from DOCX documents, preserving heading structure
and table content for comprehensive SOW review.
"""

import io
import logging
from typing import List, Tuple

from docx import Document
from docx.table import Table

logger = logging.getLogger(__name__)

# DOCX magic bytes: PK (ZIP archive)
DOCX_MAGIC_BYTES = b"PK"


class DOCXParser:
    """Extract structured text from DOCX documents using python-docx."""

    @staticmethod
    def validate_file(file_bytes: bytes) -> bool:
        """Validate that the file is a genuine DOCX by checking magic bytes.

        DOCX files are ZIP archives, so they start with PK (0x50 0x4B).

        Args:
            file_bytes: Raw file content bytes.

        Returns:
            True if the file starts with DOCX/ZIP magic bytes.
        """
        return file_bytes[:2] == DOCX_MAGIC_BYTES

    @staticmethod
    def _extract_table_text(table: Table) -> str:
        """Extract text from a DOCX table, preserving row/column structure.

        Args:
            table: A python-docx Table object.

        Returns:
            Formatted table text with pipe-delimited columns.
        """
        rows_text: List[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows_text.append(" | ".join(cells))
        return "\n".join(rows_text)

    @staticmethod
    def extract_text(file_bytes: bytes) -> str:
        """Extract full text from a DOCX file.

        Processes paragraphs (with heading markers) and tables in
        document order.

        Args:
            file_bytes: Raw DOCX file content.

        Returns:
            Extracted text with heading markers and table content.

        Raises:
            ValueError: If file is not a valid DOCX.
            RuntimeError: If text extraction fails.
        """
        if not DOCXParser.validate_file(file_bytes):
            raise ValueError("Invalid DOCX file: magic bytes do not match.")

        try:
            doc = Document(io.BytesIO(file_bytes))
            content_parts: List[str] = []

            for element in doc.element.body:
                # Handle paragraphs
                if element.tag.endswith("}p"):
                    for paragraph in doc.paragraphs:
                        if paragraph._element is element:
                            text = paragraph.text.strip()
                            if not text:
                                continue

                            # Mark headings for structure preservation
                            style_name = paragraph.style.name if paragraph.style else ""
                            if style_name.startswith("Heading"):
                                level = style_name.replace("Heading ", "").strip()
                                prefix = "#" * int(level) if level.isdigit() else "##"
                                content_parts.append(f"{prefix} {text}")
                            else:
                                content_parts.append(text)
                            break

                # Handle tables
                elif element.tag.endswith("}tbl"):
                    for table in doc.tables:
                        if table._element is element:
                            table_text = DOCXParser._extract_table_text(table)
                            if table_text.strip():
                                content_parts.append(
                                    f"\n[TABLE]\n{table_text}\n[/TABLE]\n"
                                )
                            break

            if not content_parts:
                logger.warning("DOCX contained no extractable text.")
                return ""

            full_text = "\n\n".join(content_parts)
            logger.info(
                f"Extracted text from DOCX ({len(content_parts)} elements, "
                f"{len(full_text)} characters)."
            )
            return full_text

        except Exception as e:
            logger.error(f"Failed to extract text from DOCX: {e}")
            raise RuntimeError(f"DOCX text extraction failed: {e}") from e

    @staticmethod
    def extract_sections(file_bytes: bytes) -> List[Tuple[str, str]]:
        """Extract text organized by heading sections.

        Args:
            file_bytes: Raw DOCX file content.

        Returns:
            List of (heading, section_content) tuples.
        """
        if not DOCXParser.validate_file(file_bytes):
            raise ValueError("Invalid DOCX file: magic bytes do not match.")

        doc = Document(io.BytesIO(file_bytes))
        sections: List[Tuple[str, str]] = []
        current_heading = "Introduction"
        current_content: List[str] = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = paragraph.style.name if paragraph.style else ""
            if style_name.startswith("Heading"):
                # Save previous section
                if current_content:
                    sections.append(
                        (current_heading, "\n".join(current_content))
                    )
                current_heading = text
                current_content = []
            else:
                current_content.append(text)

        # Save last section
        if current_content:
            sections.append((current_heading, "\n".join(current_content)))

        return sections
